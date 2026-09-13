# ==========================================================
#  @Crystal_Finance_Service_Bot  —  раздел P&L_Cost_Analysis
#
#  Пользователь присылает Excel с расшифровкой затрат по ЦФО,
#  бот возвращает тот же файл с колонками Score / Recommend / Комментарий.
#
#  Файлы рядом с bot.py:
#     .env          — токены и настройки
#     catalog.xlsx  — справочник статей затрат (правится в Excel)
#     prompt.md     — системный промпт
#     cache.db      — кеш разборов, создаётся автоматически
#
#  Установка:
#     pip install pyTelegramBotAPI google-genai python-dotenv openpyxl
# ==========================================================

import os
import re
import sys
import json
import time
import sqlite3
import hashlib
import tempfile
import logging
import traceback
from datetime import datetime

import telebot
from google import genai
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("pnl")
# Библиотека google-genai пишет длинное предупреждение про AFC,
# к нашему сценарию оно отношения не имеет
logging.getLogger("google_genai.models").setLevel(logging.ERROR)
logging.getLogger("httpx").setLevel(logging.WARNING)

# ---------------------------------------------------------
#  1. Настройки
# ---------------------------------------------------------
load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
MODEL = os.getenv("MODEL", "gemini-3.5-flash")
# Запасная модель на случай, если основная перегружена (ошибка 503)
FALLBACK_MODEL = os.getenv("FALLBACK_MODEL", "gemini-3.1-flash-lite")

ALLOWED_USERS = {
    int(x) for x in os.getenv("ALLOWED_USERS", "").replace(" ", "").split(",") if x
}

SHOW_COMMENT = os.getenv("SHOW_COMMENT", "да").strip().lower() in ("да", "yes", "true", "1")
BATCH_SIZE = int(os.getenv("BATCH_SIZE", "20"))

# Ограничение скорости: сколько запросов в минуту разрешено.
# На бесплатном тарифе Gemini это 10-15, поэтому по умолчанию 10.
RPM = int(os.getenv("RPM", "10"))

CATALOG_FILE = "catalog.xlsx"
PROMPT_FILE = "prompt.md"
CACHE_FILE = "cache.db"

if not BOT_TOKEN or BOT_TOKEN.startswith("вставь"):
    sys.exit("Не найден BOT_TOKEN — проверь .env")
if not GEMINI_API_KEY or GEMINI_API_KEY.startswith("вставь"):
    sys.exit("Не найден GEMINI_API_KEY — проверь .env")

bot = telebot.TeleBot(BOT_TOKEN)
client = genai.Client(api_key=GEMINI_API_KEY)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


# ---------------------------------------------------------
#  2. Справочник и промпт
# ---------------------------------------------------------
def load_catalog(path=CATALOG_FILE):
    if not os.path.exists(path):
        sys.exit(f"Не найден {path} — положи его рядом с bot.py")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Справочник"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = str(row[0]).strip() if row[0] else ""
        name = str(row[1]).strip() if row[1] else ""
        desc = str(row[2]).strip() if row[2] else ""
        use = str(row[3]).strip().lower() if len(row) > 3 and row[3] else "да"
        if name and use in ("да", "yes", "true", "1"):
            items.append({"code": code, "name": name, "desc": desc})
    wb.close()

    if not items:
        sys.exit("В catalog.xlsx нет активных статей — проверь колонку «Использовать»")
    return items


def catalog_to_text(items):
    return "\n".join(
        f"[{it['code']}] {it['name']}\n    {it['desc'] or '(описание не заполнено)'}"
        for it in items
    )


def build_prompt():
    if not os.path.exists(PROMPT_FILE):
        sys.exit(f"Не найден {PROMPT_FILE} — положи его рядом с bot.py")
    with open(PROMPT_FILE, encoding="utf-8") as f:
        template = f.read()

    items = load_catalog()
    text = catalog_to_text(items)
    version = hashlib.sha256(text.encode()).hexdigest()[:12]
    return template.replace("{catalog}", text), items, version


SYSTEM_PROMPT, CATALOG, CATALOG_VERSION = build_prompt()


# ---------------------------------------------------------
#  3. Кеш разборов (он же накопительная база исследований)
# ---------------------------------------------------------
CACHE_COLUMNS = [
    "key", "item", "content", "score", "recommend_code", "recommend_name",
    "comment", "model", "catalog_version", "created_at",
]

CREATE_SQL = """
    CREATE TABLE IF NOT EXISTS analysis (
        key             TEXT PRIMARY KEY,
        item            TEXT,
        content         TEXT,
        score           INTEGER,
        recommend_code  TEXT,
        recommend_name  TEXT,
        comment         TEXT,
        model           TEXT,
        catalog_version TEXT,
        created_at      TEXT
    )
"""


def migrate_cache(con, old_cols):
    """Структура таблицы изменилась между версиями бота.
    Переносим то, что можно перенести, остальное отбрасываем."""
    log.warning("Структура кеша устарела (%s колонок), обновляю...", len(old_cols))

    con.execute("ALTER TABLE analysis RENAME TO analysis_old")
    con.execute(CREATE_SQL)

    # переносим, только если в старой таблице есть всё необходимое
    need = {"item", "content", "score", "recommend_code",
            "recommend_name", "comment", "catalog_version"}
    moved = 0
    if need.issubset(set(old_cols)):
        rows = con.execute(
            "SELECT item, content, score, recommend_code, recommend_name, "
            "comment, model, catalog_version, created_at FROM analysis_old"
        ).fetchall()
        for r in rows:
            item, content, version = r[0], r[1], r[7]
            # ключ в новой версии считается иначе — пересчитываем
            raw = "|".join([str(version), norm(item), norm(content)])
            key = hashlib.sha256(raw.encode()).hexdigest()
            con.execute(
                f"INSERT OR REPLACE INTO analysis ({','.join(CACHE_COLUMNS)}) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (key,) + tuple(r),
            )
            moved += 1

    con.execute("DROP TABLE analysis_old")
    con.commit()
    log.warning("Кеш обновлён, перенесено записей: %s", moved)


def init_cache():
    con = sqlite3.connect(CACHE_FILE, check_same_thread=False)
    con.execute(CREATE_SQL)
    cols = [r[1] for r in con.execute("PRAGMA table_info(analysis)")]
    if cols != CACHE_COLUMNS:
        migrate_cache(con, cols)
    con.commit()
    return con


CACHE = init_cache()


def cache_key(row):
    """Ключ по сути расхода: статья + содержание + версия справочника.
    Номер документа и сумма в ключ не входят — один и тот же товар
    в разных накладных разбирается один раз."""
    raw = "|".join([CATALOG_VERSION, norm(row["item"]), norm(row["content"])])
    return hashlib.sha256(raw.encode()).hexdigest()


def cache_get(row):
    r = CACHE.execute(
        "SELECT score, recommend_code, recommend_name, comment FROM analysis WHERE key = ?",
        (cache_key(row),),
    ).fetchone()
    if not r:
        return None
    return {"score": r[0], "recommend_code": r[1], "recommend_name": r[2], "comment": r[3]}


def cache_put(row, result):
    CACHE.execute(
        f"INSERT OR REPLACE INTO analysis ({','.join(CACHE_COLUMNS)}) "
        "VALUES (?,?,?,?,?,?,?,?,?,?)",
        (
            cache_key(row), row["item"], row["content"],
            result["score"], result["recommend_code"], result["recommend_name"],
            result["comment"], MODEL, CATALOG_VERSION,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    CACHE.commit()


# ---------------------------------------------------------
#  4. Чтение файла пользователя
# ---------------------------------------------------------
COLUMN_HINTS = {
    "item": ("стать",),
    "document": ("документ",),
    "content": ("содержан",),
    "amount": ("сумма",),
}

# Итоговые строки выгрузки — это не расходы, разбирать их не надо
TOTAL_ROWS = {"итого", "всего", "итог", "total"}


def find_columns(header):
    found = {}
    for idx, title in enumerate(header):
        low = norm(title)
        for field, hints in COLUMN_HINTS.items():
            if field not in found and any(h in low for h in hints):
                found[field] = idx
    return found


def read_input(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    cols, header_idx = None, None
    for i, r in enumerate(all_rows[:20]):
        c = find_columns(r)
        if len(c) == 4:
            cols, header_idx = c, i
            break

    if cols is None:
        raise ValueError(
            "Не нашёл нужные колонки. В шапке файла должны быть слова "
            "«статья», «документ», «содержание» и «сумма»."
        )

    def cell(r, key):
        i = cols[key]
        return r[i] if i < len(r) else None

    rows, skipped_total = [], 0
    for r in all_rows[header_idx + 1:]:
        item = str(cell(r, "item") or "").strip()
        if not item:
            continue
        if norm(item) in TOTAL_ROWS:
            skipped_total += 1
            continue
        rows.append({
            "item": item,
            "document": str(cell(r, "document") or "").strip(),
            "content": str(cell(r, "content") or "").strip(),
            "amount": cell(r, "amount"),
        })

    if not rows:
        raise ValueError("Файл прочитан, но строк с данными в нём нет.")
    return rows, skipped_total


# ---------------------------------------------------------
#  5. Обращение к модели
# ---------------------------------------------------------
_last_call = 0.0
_min_interval = 60.0 / max(RPM, 1)


def throttle():
    """Держим паузу между запросами, чтобы не упереться в лимит тарифа."""
    global _last_call
    wait = _min_interval - (time.time() - _last_call)
    if wait > 0:
        time.sleep(wait)
    _last_call = time.time()


def is_rate_limit(err):
    t = str(err).lower()
    return "429" in t or "resource_exhausted" in t or "rate limit" in t or "quota" in t


def is_transient(err):
    """503/500/502/504 и таймауты — временные сбои на стороне Google.
    Лечатся простым повтором через несколько секунд."""
    t = str(err).lower()
    return any(x in t for x in (
        "503", "500", "502", "504", "unavailable", "overloaded",
        "internal error", "deadline", "timeout", "timed out",
    ))


def call_model(user_text, attempts=5, on_retry=None):
    """Запрос к модели с паузами и повторами.
    Сначала пробуем основную модель, при упорных сбоях — запасную."""
    models = [MODEL] + ([FALLBACK_MODEL] if FALLBACK_MODEL and FALLBACK_MODEL != MODEL else [])

    last_error = None
    for model_name in models:
        for n in range(attempts):
            throttle()
            try:
                return client.models.generate_content(
                    model=model_name,
                    contents=user_text,
                    config={
                        "system_instruction": SYSTEM_PROMPT,
                        "response_mime_type": "application/json",
                        "temperature": 0,
                    },
                )
            except Exception as e:
                last_error = e
                if n == attempts - 1:
                    break

                if is_rate_limit(e):
                    pause, why = 25 * (n + 1), "лимит тарифа"
                elif is_transient(e):
                    pause, why = 5 * (2 ** n), "модель перегружена"
                else:
                    raise      # настоящая ошибка — повторять бессмысленно

                log.warning("%s (%s). Пауза %s c, попытка %s из %s",
                            why, model_name, pause, n + 2, attempts)
                if on_retry:
                    on_retry(why, pause)
                time.sleep(pause)

        if model_name != models[-1]:
            log.warning("Переключаюсь на запасную модель: %s", models[-1])
            if on_retry:
                on_retry(f"переключаюсь на {models[-1]}", 0)

    raise last_error or RuntimeError("Не удалось получить ответ модели")


def parse_json(text):
    text = re.sub(r"^```(?:json)?|```$", "", (text or "").strip(), flags=re.MULTILINE).strip()
    return json.loads(text)


def analyze_batch(batch, on_retry=None):
    payload = [
        {
            "n": i + 1,
            "Статья": r["item"],
            "Документ": r["document"],
            "Содержание": r["content"],
            "Сумма": str(r["amount"]),
        }
        for i, r in enumerate(batch)
    ]
    user_text = (
        f"Проверь {len(batch)} строк расшифровки. "
        "Верни массив JSON, ровно один объект на строку.\n\n"
        + json.dumps(payload, ensure_ascii=False, indent=1)
    )

    data = parse_json(call_model(user_text, on_retry=on_retry).text)
    by_n = {int(d.get("n", 0)): d for d in data if isinstance(d, dict)}
    return [normalize(by_n.get(i, {}), row) for i, row in enumerate(batch, start=1)]


NO_CONTENT = {
    "score": None,
    "recommend_code": "",
    "recommend_name": "",
    "comment": "Нет содержания документа — по этой строке оценка невозможна.",
}


def normalize(d, row):
    try:
        score = int(d.get("score", 5))
    except (TypeError, ValueError):
        score = 5
    score = max(0, min(10, score))

    rec_code = str(d.get("recommend_code") or "").strip()
    rec_name = str(d.get("recommend_name") or "").strip()
    comment = str(d.get("comment") or "").strip()

    if not rec_name:
        return {
            "score": 5, "recommend_code": "", "recommend_name": "",
            "comment": "Модель не вернула ответ по этой строке — проверьте вручную.",
        }

    # Правило из ТЗ проверяем кодом: модель его иногда нарушает
    if norm(rec_name) != norm(row["item"]):
        score = min(score, 5)

    return {"score": score, "recommend_code": rec_code,
            "recommend_name": rec_name, "comment": comment}


# ---------------------------------------------------------
#  6. Сборка результата
# ---------------------------------------------------------
ARIAL = "Arial"
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")


def write_output(rows, results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L Cost Analysis"

    headers = ["Наименование статьи бюджета",
               "Документ (наименование, №, дата документа)",
               "Содержание документа", "Сумма", "Score", "Recommend"]
    if SHOW_COMMENT:
        headers.append("Комментарий")

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for i, (row, res) in enumerate(zip(rows, results), start=2):
        rec = f"{res['recommend_code']} {res['recommend_name']}".strip()
        values = [row["item"], row["document"], row["content"], row["amount"],
                  res["score"], rec]
        if SHOW_COMMENT:
            values.append(res["comment"])

        for c, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = Font(name=ARIAL, size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (3, 7)))

        ws.cell(row=i, column=4).number_format = "#,##0.00"
        s = ws.cell(row=i, column=5)
        s.alignment = Alignment(horizontal="center", vertical="center")
        if res["score"] is not None:
            s.fill = FILL_OK if res["score"] >= 6 else FILL_WARN if res["score"] >= 3 else FILL_BAD

    for col, w in {"A": 34, "B": 46, "C": 60, "D": 14, "E": 8, "F": 34, "G": 55}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{len(rows) + 1}"
    wb.save(path)


# ---------------------------------------------------------
#  7. Телеграм: доступ и команды
# ---------------------------------------------------------
def allowed(message):
    return message.from_user.id in ALLOWED_USERS


def deny(message):
    bot.reply_to(
        message,
        "Доступ к боту ограничен.\n"
        f"Ваш Telegram ID: {message.from_user.id}\n"
        "Передайте его администратору для подключения.",
    )


@bot.message_handler(commands=["start", "help"])
def cmd_start(message):
    if not allowed(message):
        return deny(message)
    bot.reply_to(
        message,
        "Финансовая служба — раздел P&L_Cost_Analysis\n\n"
        "Пришлите Excel с расшифровкой затрат по ЦФО.\n"
        "Нужные колонки: статья, документ, содержание, сумма.\n\n"
        "В ответ придёт тот же файл с колонками Score и Recommend.\n"
        "Строки без содержания документа не оцениваются — по ним\n"
        "судить не о чем, это автоматические проводки.\n\n"
        "/reload — перечитать справочник\n"
        "/stats — что накоплено в базе\n"
        "/whoami — узнать свой Telegram ID",
    )


@bot.message_handler(commands=["whoami"])
def cmd_whoami(message):
    bot.reply_to(message, f"Ваш Telegram ID: {message.from_user.id}")


@bot.message_handler(commands=["reload"])
def cmd_reload(message):
    if not allowed(message):
        return deny(message)
    global SYSTEM_PROMPT, CATALOG, CATALOG_VERSION
    try:
        SYSTEM_PROMPT, CATALOG, CATALOG_VERSION = build_prompt()
        bot.reply_to(message, f"Справочник обновлён: {len(CATALOG)} активных статей.")
    except Exception as e:
        bot.reply_to(message, f"Не удалось перечитать справочник: {e}")


@bot.message_handler(commands=["stats"])
def cmd_stats(message):
    if not allowed(message):
        return deny(message)
    total = CACHE.execute("SELECT COUNT(*) FROM analysis").fetchone()[0]
    low = CACHE.execute("SELECT COUNT(*) FROM analysis WHERE score <= 5").fetchone()[0]
    bot.reply_to(
        message,
        f"В базе разборов: {total} уникальных расходов\n"
        f"Из них со спорной статьёй (Score ≤ 5): {low}\n"
        f"Активных статей в справочнике: {len(CATALOG)}",
    )


# ---------------------------------------------------------
#  8. Основной сценарий: пришёл файл
# ---------------------------------------------------------
class Progress:
    """Обновляет одно и то же сообщение, но не чаще раза в 15 секунд,
    иначе Telegram начинает ругаться на частые правки."""

    def __init__(self, chat_id, message_id, period=15):
        self.chat_id, self.message_id = chat_id, message_id
        self.period, self.last, self.text = period, 0.0, ""

    def set(self, text, force=False):
        if text == self.text:
            return
        if not force and time.time() - self.last < self.period:
            return
        try:
            bot.edit_message_text(text, self.chat_id, self.message_id)
            self.text, self.last = text, time.time()
        except Exception as e:
            log.debug("не удалось обновить прогресс: %s", e)


@bot.message_handler(content_types=["document"])
def handle_document(message):
    if not allowed(message):
        return deny(message)

    name = message.document.file_name or ""
    if not name.lower().endswith((".xlsx", ".xlsm")):
        return bot.reply_to(message, "Нужен файл Excel (.xlsx).")

    status = bot.reply_to(message, "Файл получен, читаю...")
    progress = Progress(status.chat.id, status.message_id)
    tmpdir = tempfile.mkdtemp()

    try:
        info = bot.get_file(message.document.file_id)
        src = os.path.join(tmpdir, "input.xlsx")
        with open(src, "wb") as f:
            f.write(bot.download_file(info.file_path))
        rows, skipped_total = read_input(src)
    except Exception as e:
        log.error("чтение файла: %s", traceback.format_exc())
        return progress.set(f"Не смог прочитать файл.\n{e}", force=True)

    total = len(rows)
    results = [None] * total
    groups = {}          # одинаковые расходы разбираем один раз
    no_content = 0

    for i, row in enumerate(rows):
        if not row["content"]:
            results[i] = NO_CONTENT
            no_content += 1
            continue
        groups.setdefault((norm(row["item"]), norm(row["content"])), []).append(i)

    pending, from_cache = [], 0
    for idxs in groups.values():
        hit = cache_get(rows[idxs[0]])
        if hit:
            for i in idxs:
                results[i] = hit
            from_cache += len(idxs)
        else:
            pending.append(idxs)

    calls = (len(pending) + BATCH_SIZE - 1) // BATCH_SIZE
    eta = max(1, round(calls / max(RPM, 1)))

    progress.set(
        f"Строк в файле: {total}\n"
        f"Без содержания (не оцениваются): {no_content}\n"
        f"К разбору: {len(pending)} уникальных, из кеша {from_cache}\n"
        f"Запросов к модели: {calls}, примерно {eta} мин.",
        force=True,
    )

    done = 0
    try:
        for start in range(0, len(pending), BATCH_SIZE):
            chunk = pending[start:start + BATCH_SIZE]
            batch = [rows[idxs[0]] for idxs in chunk]
            def notify(why, pause):
                progress.set(
                    f"Разобрано {done} из {len(pending)}. "
                    f"{why.capitalize()}, жду {pause} c...", force=True
                )

            batch_results = analyze_batch(batch, on_retry=notify)

            for idxs, res in zip(chunk, batch_results):
                cache_put(rows[idxs[0]], res)
                for i in idxs:
                    results[i] = res

            done += len(chunk)
            progress.set(f"Разобрано {done} из {len(pending)} уникальных расходов...")
    except Exception as e:
        log.error("модель: %s", traceback.format_exc())
        hint = ""
        if is_rate_limit(e):
            hint = ("\n\nУпёрлись в лимит бесплатного тарифа. "
                    "Подождите минуту и пришлите файл снова — "
                    "разобранное уже в кеше, бот продолжит с этого места.")
        elif is_transient(e):
            hint = ("\n\nМодель перегружена на стороне Google. Это временно. "
                    "Пришлите файл через несколько минут — "
                    "разобранное уже в кеше, бот продолжит с этого места.")
        return progress.set(
            f"Ошибка после {done} из {len(pending)}.\n{str(e)[:400]}{hint}", force=True
        )

    # на всякий случай: строки, по которым почему-то нет результата
    for i, r in enumerate(results):
        if r is None:
            results[i] = NO_CONTENT

    out = os.path.join(tmpdir, f"P&L_analysis_{datetime.now():%Y-%m-%d_%H%M}.xlsx")
    write_output(rows, results, out)

    scored = [r["score"] for r in results if r["score"] is not None]
    low = sum(1 for s in scored if s <= 5)
    avg = sum(scored) / len(scored) if scored else 0

    progress.set(f"Готово. Обработано {total} строк.", force=True)
    with open(out, "rb") as f:
        bot.send_document(
            message.chat.id, f,
            caption=(
                f"Всего строк: {total}\n"
                f"Оценено: {len(scored)}\n"
                f"Спорных (Score ≤ 5): {low}\n"
                f"Средний Score: {avg:.1f}\n"
                f"Без содержания: {no_content}"
                + (f"\nИтоговых строк пропущено: {skipped_total}" if skipped_total else "")
            ),
        )


@bot.message_handler(func=lambda m: True, content_types=["text"])
def fallback(message):
    if not allowed(message):
        return deny(message)
    bot.reply_to(message, "Пришлите Excel-файл с расшифровкой затрат. Подсказка: /help")


# ---------------------------------------------------------
#  9. Запуск
# ---------------------------------------------------------
if __name__ == "__main__":
    log.info("Справочник: %s активных статей (версия %s)", len(CATALOG), CATALOG_VERSION)
    log.info("Доступ разрешён: %s пользователям", len(ALLOWED_USERS))
    log.info("Модель %s (запасная %s), не более %s запросов в минуту, пачка %s",
             MODEL, FALLBACK_MODEL, RPM, BATCH_SIZE)
    log.info("Бот запущен.")

    # Разбор большого файла занимает минуты, за это время связь с Telegram
    # может оборваться. Поднимаемся сами, вручную перезапускать не нужно.
    while True:
        try:
            bot.infinity_polling(timeout=30, long_polling_timeout=30, skip_pending=True)
            log.warning("Опрос Telegram завершился, перезапускаю через 5 с...")
        except Exception:
            log.error("Сбой опроса Telegram:\n%s", traceback.format_exc())
        time.sleep(5)
